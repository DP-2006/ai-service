import io
import json
import re

from app.services.ollama_client import chat_text


class PresentationService:
    def build_slides(self, topic: str, context: str, slide_count: int = 8, system: str = "") -> list:
        prompt = (
            f"TOPIC: {topic}\n"
            f"SLIDE_COUNT: {slide_count}\n\n"
            f"SOURCES:\n{context}\n\n"
            f"Return strict JSON: {{\"slides\":[{{\"title\":str,\"bullets\":[str],\"notes\":str}}]}}"
        )
        raw = chat_text(prompt, system=system)
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if not match:
            return [{"title": topic, "bullets": ["(no content)"], "notes": ""}]
        try:
            data = json.loads(match.group(0))
            return data.get("slides", [])
        except json.JSONDecodeError:
            return [{"title": topic, "bullets": ["(invalid JSON)"], "notes": raw[:500]}]

    def render_pptx(self, title: str, slides: list, rtl: bool = True) -> bytes:
        try:
            from pptx import Presentation
            from pptx.util import Pt
        except ImportError:
            raise RuntimeError("python-pptx not installed. Run: pip install python-pptx")

        prs = Presentation()
        # slide 1
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
        # slides
        for sl in slides:
            layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(layout)
            s.shapes.title.text = sl.get("title", "")
            body = s.placeholders[1].text_frame
            body.clear()
            for i, b in enumerate(sl.get("bullets", [])):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = b
                p.font.size = Pt(18)

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
